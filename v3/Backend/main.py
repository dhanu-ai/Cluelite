# main.py — v3 with mode switching and fixed ping/root routes
import os
import threading
import logging
import uuid
from fastapi import FastAPI, Body, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from dotenv import load_dotenv
from docx import Document
from typing import Dict, Any
from datetime import datetime
import time

# Load environment variables
load_dotenv()

# Logging setup
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s"
)
logger = logging.getLogger("Cluelite")

# Core modules
from ocr import ocr_loop, get_latest_ocr_text, clear_ocr_history
from stt import get_latest_transcripts, start_stt_streams, stop_stt_streams, clear_transcripts
from groq_ai import generate_suggestion
from suggestion_loop import run_suggestion_loop, get_latest_suggestion

# FastAPI App
app = FastAPI(
    title="Cluelite API",
    description="AI-powered interview assistant backend (v3)",
    version="3.0.0"
)

# CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=os.getenv("ALLOWED_ORIGINS", "*").split(","),
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["*"],
    allow_credentials=True,
)

# Global state
stop_flag = threading.Event()
threads_started = False
active_sessions: Dict[str, Any] = {}


# === Utility: Generate DOCX Summary ===
def generate_docx_summary(transcript_text: str, ocr_lines: list, analysis_text: str, file_path: str) -> str:
    try:
        doc = Document()
        doc.add_heading("Cluelite Session Summary", 0)
        doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(f"Session ID: {uuid.uuid4()}")
        doc.add_page_break()

        # Transcripts
        doc.add_heading("1. Conversation Transcript", 1)
        doc.add_paragraph(transcript_text.strip() or "No transcript available.")

        # Screen content
        doc.add_heading("2. Screen Content", 1)
        if ocr_lines:
            for line in ocr_lines:
                doc.add_paragraph(f"- {line}")
        else:
            doc.add_paragraph("No screen content captured.")

        # Analysis
        doc.add_heading("3. Performance Analysis & Recommendations", 1)
        doc.add_paragraph(analysis_text.strip() or "No analysis available.")

        # Save with retries
        attempt = 0
        while attempt < 3:
            try:
                doc.save(file_path)
                break
            except PermissionError:
                attempt += 1
                logger.warning("DOCX save PermissionError; retrying...")
                time.sleep(0.3)
        if not os.path.exists(file_path):
            raise RuntimeError("DOCX file was not created")

        logger.info(f"DOCX summary generated: {file_path}")
        return file_path
    except Exception as e:
        logger.error(f"Error generating DOCX summary: {e}")
        raise HTTPException(status_code=500, detail="Failed to generate summary document")


# === Routes ===
@app.get("/")
async def root():
    return {
        "message": "Cluelite API Server (v3)",
        "status": "healthy",
        "version": "3.0.0"
    }


@app.get("/ping")
def ping():
    return {
        "status": "alive",
        "timestamp": datetime.now().isoformat()
    }


@app.post("/start")
def start_monitoring():
    global threads_started, stop_flag
    if threads_started:
        return {"status": "already_running", "message": "Monitoring is already active"}

    stop_flag.clear()
    clear_transcripts()
    clear_ocr_history()

    threading.Thread(target=ocr_loop, args=(stop_flag,), daemon=True).start()
    logger.info("✅ OCR thread started")

    threading.Thread(target=run_suggestion_loop, args=(stop_flag,), daemon=True).start()
    logger.info("✅ Suggestion loop started")

    start_stt_streams()
    logger.info("🎙️ STT streams started")

    threads_started = True
    session_id = str(uuid.uuid4())
    active_sessions[session_id] = {"start_time": datetime.now().isoformat(), "status": "active"}
    logger.info("📈 Monitoring started")
    return {"status": "started", "session_id": session_id}


@app.get("/suggestion")
def api_get_suggestion():
    try:
        return {"suggestion": get_latest_suggestion()}
    except Exception as e:
        logger.error(f"Error serving suggestion: {e}")
        raise HTTPException(status_code=500, detail="Failed to get suggestion")


@app.post("/askcluelite")
def ask_cluelite(query: str = Body(..., embed=True)):
    if not query.strip():
        return {"response": "Please provide a query."}
    try:
        transcripts = get_latest_transcripts()
        ocr_lines = get_latest_ocr_text()
        ocr_text = " | ".join(ocr_lines)
        combined_context = (
            f"{query}\n\n"
            f"[SCREEN CONTEXT]: {ocr_text}\n"
            f"[TRANSCRIPT CONTEXT]: {transcripts.get('interviewee','')} {transcripts.get('interviewer','')}"
        )

        response = generate_suggestion(
            transcript=combined_context,
            ocr_text=ocr_text,
            mode=os.getenv("GROQ_ASK_MODE", "response"),
            timeout=8
        )
        return {"response": response or "I couldn't generate a response. Please try again."}
    except Exception as e:
        logger.error(f"Error processing query: {e}")
        raise HTTPException(status_code=500, detail="Failed to process query")


@app.post("/stop")
def stop_monitoring():
    global threads_started, stop_flag
    try:
        final_ocr = get_latest_ocr_text()
        final_transcripts = get_latest_transcripts()
        full_transcript = (
            f"INTERVIEWER:\n{final_transcripts.get('interviewer','')}\n\n"
            f"INTERVIEWEE:\n{final_transcripts.get('interviewee','')}"
        )

        if not final_transcripts.get('interviewer') and not final_transcripts.get('interviewee'):
            full_transcript = "Note: Limited audio captured during this session\n\n" + full_transcript

        analysis = generate_suggestion(
            transcript=full_transcript,
            ocr_text=" | ".join(final_ocr),
            mode=os.getenv("GROQ_ANALYSIS_MODE", "analysis"),
            timeout=12
        )

        stop_flag.set()
        stop_stt_streams()
        threads_started = False

        session_id = str(uuid.uuid4())
        docx_file = f"cluelite_summary_{session_id}.docx"
        docx_path = os.path.join(os.getcwd(), docx_file)
        generate_docx_summary(full_transcript, final_ocr, analysis, docx_path)

        return {
            "status": "Session ended successfully",
            "session_id": session_id,
            "docx_file": docx_file,
            "download_url": f"/download/{docx_file}"
        }
    except Exception as e:
        logger.error(f"Error stopping monitoring: {e}")
        raise HTTPException(status_code=500, detail="Failed to stop monitoring session")


@app.get("/download/{filename}")
async def download_file(filename: str):
    if not filename.endswith('.docx') or '..' in filename or '/' in filename:
        raise HTTPException(status_code=400, detail="Invalid filename")
    file_path = os.path.join(os.getcwd(), filename)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(
        file_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="cluelite_session_report.docx"
    )


@app.get("/health")
def health_check():
    return {
        "status": "healthy" if threads_started else "idle",
        "threads_started": threads_started,
        "active_sessions": len(active_sessions),
        "timestamp": datetime.now().isoformat()
    }


@app.exception_handler(Exception)
async def global_exception_handler(request, exc):
    logger.error(f"Unhandled exception: {exc}")
    return JSONResponse(status_code=500, content={"message": "An internal server error occurred"})
