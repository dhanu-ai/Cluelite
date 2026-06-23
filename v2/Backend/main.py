# main.py
import os
import threading
import logging
import uuid
from fastapi import FastAPI, Body, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from dotenv import load_dotenv
from docx import Document
from typing import Dict, Any
from datetime import datetime
import time
from starlette.responses import JSONResponse

# Load environment variables
load_dotenv()

# === Logging ===
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s"
)
logger = logging.getLogger("Cluelite")

# === Core Modules ===
from ocr import ocr_loop, get_latest_ocr_text, clear_ocr_history, pause_ocr_processing, resume_ocr_processing
from stt import get_latest_transcripts, start_stt_streams, stop_stt_streams, clear_transcripts, pause_processing, resume_processing
from groq_ai import generate_suggestion
from suggestion_loop import run_suggestion_loop, get_latest_suggestion

# === FastAPI App ===
app = FastAPI(
    title="Cluelite API",
    description="AI-powered interview assistant backend",
    version="1.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=os.getenv("ALLOWED_ORIGINS", "*").split(","),
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["*"],
    allow_credentials=True,
)

# === Global State ===
stop_flag = threading.Event()
threads_started = False
active_sessions: Dict[str, Any] = {}
system_load = "normal"  # normal, high, critical

# === DOCX Generator ===
def generate_docx_summary(transcript_text: str, ocr_lines: list, analysis_text: str, file_path: str) -> str:
    """Generate a professional DOCX summary report."""
    try:
        doc = Document()
        
        # Title page
        doc.add_heading("Cluelite Session Summary", 0)
        doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(f"Session ID: {uuid.uuid4()}")
        doc.add_page_break()
        
        # Transcripts section
        doc.add_heading("1. Conversation Transcript", 1)
        if transcript_text.strip():
            doc.add_paragraph(transcript_text)
        else:
            doc.add_paragraph("No transcript available for this session.")
        
        # Screen content section
        doc.add_heading("2. Screen Content", 1)
        if ocr_lines:
            doc.add_paragraph("\n".join([f"- {line}" for line in ocr_lines]))
        else:
            doc.add_paragraph("No screen content captured.")
        
        # Analysis section
        doc.add_heading("3. Performance Analysis & Recommendations", 1)
        if analysis_text.strip():
            doc.add_paragraph(analysis_text)
        else:
            doc.add_paragraph("No analysis available for this session.")
        
        # Save document
        doc.save(file_path)
        logger.info(f"DOCX summary generated: {file_path}")
        return file_path
        
    except Exception as e:
        logger.error(f"Error generating DOCX summary: {e}")
        raise HTTPException(status_code=500, detail="Failed to generate summary document")

# === Resource Management ===
def adjust_system_load():
    """Adjust system processing based on current load."""
    global system_load
    
    # Check if we need to reduce processing load
    if system_load == "high":
        pause_ocr_processing()
        logger.warning("System load high, pausing OCR processing")
    elif system_load == "critical":
        pause_processing()
        pause_ocr_processing()
        logger.error("System load critical, pausing all processing")
    else:
        resume_processing()
        resume_ocr_processing()
        logger.info("System load normal, resuming all processing")

# === Routes ===

@app.get("/")
async def root():
    return {"message": "Cluelite API Server", "status": "healthy", "version": "1.0.0"}

@app.get("/ping")
def ping():
    return {"status": "alive", "timestamp": datetime.now().isoformat()}

@app.post("/start")
def start_monitoring():
    global threads_started, stop_flag, system_load
    if not threads_started:
        stop_flag.clear()
        clear_transcripts()
        clear_ocr_history()
        system_load = "normal"

        # Start all processing threads
        threading.Thread(target=ocr_loop, args=(stop_flag,), daemon=True, name="ocr_thread").start()
        logger.info("✅ OCR thread started")

        threading.Thread(target=run_suggestion_loop, args=(stop_flag,), daemon=True, name="suggestion_thread").start()
        logger.info("✅ Suggestion loop started")

        if start_stt_streams():
            logger.info("🎙️ STT streams started")
        else:
            logger.error("❌ Failed to start STT streams")

        threads_started = True
        session_id = str(uuid.uuid4())
        active_sessions[session_id] = {
            "start_time": datetime.now().isoformat(),
            "status": "active"
        }
        logger.info("📈 Monitoring started")
        return {"status": "started", "session_id": session_id}
    return {"status": "already_running", "message": "Monitoring is already active"}

@app.get("/suggestion")
def get_suggestion():
    try:
        transcripts = get_latest_transcripts()
        ocr_lines = get_latest_ocr_text()
        combined_transcript = f"{transcripts.get('interviewee', '')} {transcripts.get('interviewer', '')}"

        if not combined_transcript.strip() and not any(ocr_lines):
            return {"suggestion": "Waiting for conversation or screen content..."}

        suggestion = generate_suggestion(
            transcript=combined_transcript,
            ocr_text=" | ".join(ocr_lines),
            mode="auto",
            timeout=5
        )

        return {"suggestion": suggestion or "Analyzing content, please wait..."}
    
    except Exception as e:
        logger.error(f"Error generating suggestion: {e}")
        raise HTTPException(status_code=500, detail="Failed to generate suggestion")

@app.post("/askcluelite")
def ask_cluelite(query: str = Body(..., embed=True)):
    if not query.strip():
        return {"response": "Please provide a query."}

    try:
        transcripts = get_latest_transcripts()
        ocr_lines = get_latest_ocr_text()
        ocr_text = " | ".join(ocr_lines)

        combined_context = f"{query}\n\n[SCREEN CONTEXT]: {ocr_text}\n[TRANSCRIPT CONTEXT]: {transcripts.get('interviewee', '')} {transcripts.get('interviewer', '')}"

        response = generate_suggestion(
            transcript=combined_context,
            ocr_text=ocr_text,
            mode="response",
            timeout=5
        )

        return {"response": response or "I couldn't generate a response. Please try again."}
    
    except Exception as e:
        logger.error(f"Error processing query: {e}")
        raise HTTPException(status_code=500, detail="Failed to process query")

@app.post("/stop")
def stop_monitoring():
    global threads_started, stop_flag

    try:
        final_ocr = get_latest_ocr_text()
        final_transcripts = get_latest_transcripts()
        full_transcript = f"INTERVIEWER:\n{final_transcripts.get('interviewer', '')}\n\nINTERVIEWEE:\n{final_transcripts.get('interviewee', '')}"

        if not final_transcripts.get('interviewer', '') and not final_transcripts.get('interviewee', ''):
            full_transcript = "Note: Limited audio captured during this session\n\n" + full_transcript

        # Generate AI analysis
        analysis = generate_suggestion(
            transcript=full_transcript,
            ocr_text=" | ".join(final_ocr),
            mode="analysis",
            timeout=10
        )

        stop_flag.set()
        stop_stt_streams()
        threads_started = False

        # Create DOCX summary
        session_id = str(uuid.uuid4())
        docx_file = f"cluelite_summary_{session_id}.docx"
        docx_path = generate_docx_summary(full_transcript, final_ocr, analysis, docx_file)

        return {
            "status": "Session ended successfully", 
            "session_id": session_id,
            "docx_file": docx_file,
            "download_url": f"/download/{docx_file}"
        }
    
    except Exception as e:
        logger.error(f"Error stopping monitoring: {e}")
        raise HTTPException(status_code=500, detail="Failed to stop monitoring session")

@app.get("/download/{filename}")
async def download_file(filename: str):
    if not filename.endswith('.docx') or '..' in filename or '/' in filename:
        raise HTTPException(status_code=400, detail="Invalid filename")
    
    file_path = os.path.join(os.getcwd(), filename)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File not found")
    
    return FileResponse(
        file_path, 
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"cluelite_session_report.docx"
    )

@app.get("/health")
def health_check():
    return {
        "status": "healthy",
        "threads_started": threads_started,
        "active_sessions": len(active_sessions),
        "system_load": system_load,
        "timestamp": datetime.now().isoformat()
    }

@app.post("/adjust-load/{load_level}")
def adjust_load(load_level: str):
    global system_load
    if load_level in ["normal", "high", "critical"]:
        system_load = load_level
        adjust_system_load()
        return {"status": "success", "system_load": system_load}
    else:
        return {"status": "error", "message": "Invalid load level"}

@app.exception_handler(Exception)
async def global_exception_handler(request, exc):
    logger.error(f"Unhandled exception: {exc}")
    return JSONResponse(
        status_code=500,
        content={"message": "An internal server error occurred"},
    )